"""Import brand price-list items from bundled JSON or Excel files."""

from __future__ import annotations

import json
from pathlib import Path

import frappe
from frappe.utils import flt

DATA_FILE = Path(__file__).resolve().parent / "data" / "price_list_items.json"
DEFAULT_STOCK_UOM = "Nos"
SELLING_PRICE_LIST = "SACO Selling"


def run(
	xlsx_path: str | None = None,
	brand: str | None = None,
	sheet_name: str | None = None,
	company: str | None = None,
	update_existing: bool = True,
):
	"""Import items + base selling prices.

	Bundled data (FIABILA + YC):
	  bench --site SITE execute saco_management.setup.import_items.run

	Future brand from Excel:
	  bench --site SITE execute saco_management.setup.import_items.run \\
	    --kwargs '{"xlsx_path": "/path/PRICE LIST BRAND.xlsx", "brand": "BRANDNAME"}'
	"""
	frappe.only_for("System Manager")
	company = company or frappe.defaults.get_global_default("company")
	if not company:
		frappe.throw("Set a default Company before importing items.")

	if xlsx_path:
		rows = _parse_xlsx(xlsx_path, brand=brand, sheet_name=sheet_name)
	else:
		rows = _load_bundled_rows()

	ensure_brand_groups({row["brand"] for row in rows})
	price_list = ensure_selling_price_list(company)

	created = updated = skipped = barcode_skipped = 0
	used_barcodes: set[str] = set()
	for row in rows:
		action, bc_skipped = _upsert_item(
			row, company=company, update_existing=update_existing, used_barcodes=used_barcodes
		)
		_upsert_item_price(row, price_list, company)
		if bc_skipped:
			barcode_skipped += 1
		if action == "created":
			created += 1
		elif action == "updated":
			updated += 1
		else:
			skipped += 1

	frappe.db.commit()
	result = {
		"created": created,
		"updated": updated,
		"skipped": skipped,
		"barcode_skipped": barcode_skipped,
		"total": len(rows),
	}
	print(result)  # visible in bench execute output on VPS
	return result


def _load_bundled_rows() -> list[dict]:
	if not DATA_FILE.is_file():
		frappe.throw(
			f"Missing bundled import data: {DATA_FILE}. "
			"Push the latest saco_erp repo and rebuild the Docker image."
		)
	rows = json.loads(DATA_FILE.read_text(encoding="utf-8"))
	rows = _dedupe_rows(rows)
	print(f"Loaded {len(rows)} rows from {DATA_FILE}")
	return rows


def _dedupe_rows(rows: list[dict]) -> list[dict]:
	seen: set[str] = set()
	deduped: list[dict] = []
	for row in rows:
		code = row.get("item_code")
		if not code or code in seen:
			continue
		seen.add(code)
		deduped.append(row)
	return deduped


def _parse_xlsx(xlsx_path: str, brand: str | None = None, sheet_name: str | None = None) -> list[dict]:
	try:
		import openpyxl
	except ImportError:
		frappe.throw("Install openpyxl on the bench env to import from Excel.")

	path = Path(xlsx_path).expanduser()
	if not path.is_file():
		frappe.throw(f"Excel file not found: {path}")

	wb = openpyxl.load_workbook(path, data_only=True)
	ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
	brand_name = (brand or _guess_brand_from_sheet(ws)).strip()
	if not brand_name:
		frappe.throw("Pass brand=BRANDNAME when importing a new Excel file.")

	headers = {}
	for col in range(1, ws.max_column + 1):
		value = ws.cell(3, col).value
		if value:
			headers[str(value).strip().upper()] = col

	code_col = headers.get("CODE")
	desc_col = headers.get("DESCRIPTION")
	price_col = headers.get("PRICE")
	barcode_col = headers.get("BAR CODE") or headers.get("BARCODE") or headers.get("BARCODES")
	nw_col = headers.get("N.W") or headers.get("N.WT.")
	ctn_col = headers.get("CTN PACKING")

	if not all([code_col, desc_col, price_col]):
		frappe.throw("Excel row 3 must include CODE, DESCRIPTION, and PRICE columns.")

	rows = []
	for row_idx in range(4, ws.max_row + 1):
		code = ws.cell(row_idx, code_col).value
		desc = ws.cell(row_idx, desc_col).value
		if not code or not desc:
			continue
		code = str(code).strip()
		if code.upper() == "CODE":
			continue

		barcode_val = ws.cell(row_idx, barcode_col).value if barcode_col else ""
		barcode = ""
		if barcode_val:
			barcode = str(int(barcode_val)) if isinstance(barcode_val, (int, float)) else str(barcode_val).strip()

		ctn_val = ws.cell(row_idx, ctn_col).value if ctn_col else ""
		carton = ""
		if ctn_val not in (None, ""):
			carton = str(int(ctn_val)) if isinstance(ctn_val, (int, float)) else str(ctn_val).strip()

		nw_val = ws.cell(row_idx, nw_col).value if nw_col else ""
		net_weight = str(nw_val).strip() if nw_val not in (None, "") else ""

		rows.append(
			{
				"brand": brand_name,
				"item_code": code,
				"item_name": str(desc).strip(),
				"net_weight": net_weight,
				"carton_packing": carton,
				"price": flt(ws.cell(row_idx, price_col).value),
				"barcode": barcode,
			}
		)

	if not rows:
		frappe.throw(f"No item rows found in {path}")
	return rows


def _guess_brand_from_sheet(ws) -> str:
	title = str(ws.cell(1, 1).value or "")
	for token in ("FIABILA", "YC"):
		if token in title.upper():
			return token
	if "PRICE LIST OF" in title.upper():
		return title.upper().replace("PRICE LIST OF", "").replace("BRAND", "").strip(" *")
	return ""


def ensure_brand_groups(brands: set[str]) -> None:
	parent = "Products" if frappe.db.exists("Item Group", "Products") else "All Item Groups"
	if not frappe.db.exists("Item Group", parent):
		parent = "All Item Groups"
	for brand in sorted(brands):
		if frappe.db.exists("Item Group", brand):
			continue
		doc = frappe.get_doc(
			{
				"doctype": "Item Group",
				"item_group_name": brand,
				"parent_item_group": parent,
				"is_group": 0,
			}
		)
		doc.insert(ignore_permissions=True)


def ensure_selling_price_list(company: str) -> str:
	currency = frappe.get_cached_value("Company", company, "default_currency") or "PKR"
	if not frappe.db.exists("Price List", SELLING_PRICE_LIST):
		doc = frappe.get_doc(
			{
				"doctype": "Price List",
				"price_list_name": SELLING_PRICE_LIST,
				"currency": currency,
				"buying": 0,
				"selling": 1,
				"enabled": 1,
			}
		)
		doc.insert(ignore_permissions=True)

	selling = frappe.get_single("Selling Settings")
	if selling.selling_price_list != SELLING_PRICE_LIST:
		selling.selling_price_list = SELLING_PRICE_LIST
		selling.save(ignore_permissions=True)
	return SELLING_PRICE_LIST


def _has_custom_field(fieldname: str) -> bool:
	return bool(frappe.db.exists("Custom Field", {"dt": "Item", "fieldname": fieldname}))


def _barcode_available(barcode: str, item_code: str, used_barcodes: set[str]) -> bool:
	if not barcode:
		return False
	if barcode in used_barcodes:
		return False
	existing_item = frappe.db.get_value("Item Barcode", {"barcode": barcode}, "parent")
	if existing_item and existing_item != item_code:
		return False
	return True


def _upsert_item(
	row: dict, company: str, update_existing: bool, used_barcodes: set[str]
) -> tuple[str, bool]:
	item_code = row["item_code"]
	fields = {
		"item_name": row["item_name"],
		"item_group": row["brand"],
		"stock_uom": DEFAULT_STOCK_UOM,
		"is_stock_item": 1,
		"include_item_in_manufacturing": 0,
		"description": row["item_name"],
	}
	if _has_custom_field("saco_brand"):
		fields["saco_brand"] = row["brand"]
	if _has_custom_field("net_weight"):
		fields["net_weight"] = row.get("net_weight") or ""
	if _has_custom_field("carton_packing"):
		fields["carton_packing"] = row.get("carton_packing") or ""

	barcode = (row.get("barcode") or "").strip()
	barcode_skipped = bool(barcode and not _barcode_available(barcode, item_code, used_barcodes))

	if frappe.db.exists("Item", item_code):
		if not update_existing:
			return "skipped", barcode_skipped
		item = frappe.get_doc("Item", item_code)
		for key, value in fields.items():
			item.set(key, value)
		_set_barcode(item, barcode, item_code, used_barcodes)
		item.save(ignore_permissions=True)
		return "updated", barcode_skipped

	item = frappe.get_doc({"doctype": "Item", "item_code": item_code, **fields})
	_set_barcode(item, barcode, item_code, used_barcodes)
	item.insert(ignore_permissions=True)
	return "created", barcode_skipped


def _set_barcode(item, barcode: str | None, item_code: str, used_barcodes: set[str]) -> None:
	if not _barcode_available(barcode or "", item_code, used_barcodes):
		return
	item.barcodes = []
	item.append("barcodes", {"barcode": barcode})
	used_barcodes.add(barcode)


def _upsert_item_price(row: dict, price_list: str, company: str) -> None:
	item_code = row["item_code"]
	rate = flt(row.get("price"))
	name = frappe.db.get_value(
		"Item Price",
		{"item_code": item_code, "price_list": price_list},
		"name",
	)
	if name:
		doc = frappe.get_doc("Item Price", name)
		doc.price_list_rate = rate
		doc.save(ignore_permissions=True)
		return

	doc = frappe.get_doc(
		{
			"doctype": "Item Price",
			"item_code": item_code,
			"price_list": price_list,
			"price_list_rate": rate,
			"selling": 1,
			"currency": frappe.get_cached_value("Price List", price_list, "currency"),
		}
	)
	doc.insert(ignore_permissions=True)
